from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
for name, size in [('Normal', 11), ('Title', 26), ('Heading 1', 17), ('Heading 2', 13)]:
    style = doc.styles[name]
    style.font.name = 'Arial Unicode MS'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    style.paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.17
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.core_properties.title = 'AI员工工作台产品需求文档'
doc.core_properties.author = 'My Workbench'

def text_runs(p, text):
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\))', text):
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
        if not match:
            p.add_run(part)
            continue
        rel = p.part.relate_to(match[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '274EB4'); props.append(color)
        run.append(props); node = OxmlElement('w:t'); node.text = match[1]; run.append(node); link.append(run)
        p._p.append(link)

def table(rows):
    n = len(rows[0]); t = doc.add_table(rows=1, cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = {2: [1.65, 5.35], 3: [1.15, 3.65, 2.2], 4: [0.65, 0.65, 1.65, 4.05]}[n]
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (c, text) in enumerate(zip(cells, row)):
            c.width = Inches(widths[j]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.08
            if j == 0 or (n == 4 and j == 1): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text); r.font.size = Pt(10)
            props = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for tag in ['top', 'left', 'bottom', 'right']:
                x = OxmlElement('w:' + tag); x.set(qn('w:w'), '85'); x.set(qn('w:type'), 'dxa'); margins.append(x)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for tag in ['top', 'left', 'bottom', 'right']:
                x = OxmlElement('w:' + tag); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:color'), 'D9D9D9'); borders.append(x)
            props.append(borders)
            if i == 0 or i % 2 == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF0F8' if i == 0 else 'F6F8FA'); props.append(shade)
            if i == 0: r.bold = True
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = (ROOT / 'docs/PRD.md').read_text().splitlines(); i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line: i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            row = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', x) for x in row): rows.append(row)
            i += 1
        table(rows); continue
    if line.startswith('# '): doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    else: text_runs(doc.add_paragraph(), line)
    i += 1
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('My Workbench  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor(90, 90, 90)
doc.save(ROOT / 'docs/AI员工工作台_PRD_v0.3.docx')
print('PRD DOCX created')
